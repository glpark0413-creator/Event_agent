#!/usr/bin/env python3
"""
이벤트 데이터 읽기 스크립트 (Phase 1)

Usage:
  xlsx:  python read_sheet.py --file <path_to_xlsx>
  URL:   python read_sheet.py --url <spreadsheet_url>   (Google API 인증 필요)

Output: output/raw_sheet_data.json
"""

import argparse
import json
import re
import sys
from pathlib import Path


# ──────────────────────────────────────────
# xlsx 읽기
# ──────────────────────────────────────────

def read_xlsx(file_path: str) -> dict:
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl이 없습니다. 다음 명령어로 설치하세요:")
        print("  pip install openpyxl")
        sys.exit(1)

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    print(f"  시트: '{ws.title}'")

    # 헤더 (첫 번째 행)
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value) if cell.value is not None else '')

    # 데이터 행
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        row_values = []
        for cell in row:
            val = cell.value
            if val is None:
                row_values.append('')
            elif hasattr(val, 'strftime'):
                row_values.append(val.strftime('%Y/%m/%d'))
            else:
                row_values.append(str(val))
        padded = row_values + [''] * (len(headers) - len(row_values))
        rows.append(padded[:len(headers)])

    # 완전히 빈 행 제거
    rows = [r for r in rows if any(v.strip() != '' for v in r)]

    format_info = _extract_xlsx_format_info(ws, len(headers))

    return {
        "spreadsheet_id": path.stem,
        "spreadsheet_url": None,
        "source_file": str(path.resolve()),
        "sheet_name": ws.title,
        "sheet_id": 0,
        "headers": headers,
        "rows": rows,
        "total_rows": len(rows),
        "format_info": format_info
    }


def _extract_xlsx_format_info(ws, num_columns: int) -> dict:
    from openpyxl.utils import get_column_letter

    format_info = {
        "column_widths": [],
        "frozen_rows": 0,
        "header_background_color": None,
        "header_text_bold": False,
        "header_text_color": None,
    }

    # 열 너비 (엑셀 너비 단위 → 픽셀 근사: *7)
    for i in range(num_columns):
        col_letter = get_column_letter(i + 1)
        dim = ws.column_dimensions.get(col_letter)
        width_px = int((dim.width if dim and dim.width else 8.43) * 7)
        format_info["column_widths"].append(width_px)

    # 고정 행
    if ws.freeze_panes:
        match = re.match(r'[A-Z]+(\d+)', str(ws.freeze_panes))
        if match:
            format_info["frozen_rows"] = int(match.group(1)) - 1

    # 헤더 서식 (첫 번째 행 첫 번째 셀)
    first_cell = ws.cell(row=1, column=1)
    try:
        if first_cell.fill and first_cell.fill.fgColor:
            argb = first_cell.fill.fgColor.rgb
            if argb and argb not in ('00000000', '000000'):
                rgb = argb[-6:]
                format_info["header_background_color"] = {
                    "red": int(rgb[0:2], 16) / 255,
                    "green": int(rgb[2:4], 16) / 255,
                    "blue": int(rgb[4:6], 16) / 255
                }
    except Exception:
        pass

    try:
        if first_cell.font:
            format_info["header_text_bold"] = bool(first_cell.font.bold)
            if first_cell.font.color and first_cell.font.color.type == 'rgb':
                rgb = first_cell.font.color.rgb[-6:]
                format_info["header_text_color"] = {
                    "red": int(rgb[0:2], 16) / 255,
                    "green": int(rgb[2:4], 16) / 255,
                    "blue": int(rgb[4:6], 16) / 255
                }
    except Exception:
        pass

    return format_info


# ──────────────────────────────────────────
# Google Sheets 읽기 (--url 옵션)
# ──────────────────────────────────────────

def get_spreadsheet_id(url: str) -> str:
    match = re.search(r'/spreadsheets/d/([a-zA-Z0-9-_]+)', url)
    if not match:
        raise ValueError(f"유효한 Google Sheets URL이 아닙니다: {url}")
    return match.group(1)


def get_sheets_service():
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
    except ImportError:
        print("ERROR: 필요한 패키지가 없습니다. pip install -r requirements.txt")
        sys.exit(1)

    SCOPES = ['https://www.googleapis.com/auth/spreadsheets']
    creds = None
    token_path = Path('token.json')
    creds_path = Path('credentials.json')
    sa_path = Path('service_account.json')

    if sa_path.exists():
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_file(
            str(sa_path), scopes=SCOPES
        )
        return build('sheets', 'v4', credentials=creds)

    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not creds_path.exists():
                print("ERROR: credentials.json 또는 service_account.json 파일이 없습니다.")
                sys.exit(1)
            flow = InstalledAppFlow.from_client_secrets_file(str(creds_path), SCOPES)
            creds = flow.run_local_server(port=0)
        with open(token_path, 'w', encoding='utf-8') as f:
            f.write(creds.to_json())

    return build('sheets', 'v4', credentials=creds)


def read_sheet_url(url: str) -> dict:
    service = get_sheets_service()
    spreadsheet_id = get_spreadsheet_id(url)

    print(f"  Sheets 메타데이터 읽는 중...")
    spreadsheet = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        includeGridData=True
    ).execute()

    sheets = spreadsheet.get('sheets', [])
    if not sheets:
        raise ValueError("시트를 찾을 수 없습니다.")

    first_sheet = sheets[0]
    sheet_name = first_sheet['properties']['title']
    sheet_id = first_sheet['properties']['sheetId']
    print(f"  첫 번째 시트: '{sheet_name}'")

    print(f"  데이터 읽는 중...")
    result = service.spreadsheets().values().get(
        spreadsheetId=spreadsheet_id,
        range=sheet_name,
        valueRenderOption='FORMATTED_VALUE',
        dateTimeRenderOption='FORMATTED_STRING'
    ).execute()

    values = result.get('values', [])
    if not values:
        raise ValueError("시트에 데이터가 없습니다.")

    headers = values[0] if values else []
    rows = values[1:] if len(values) > 1 else []

    normalized_rows = []
    for row in rows:
        padded = row + [''] * (len(headers) - len(row))
        normalized_rows.append(padded[:len(headers)])

    format_info = _extract_gsheets_format_info(first_sheet)

    return {
        "spreadsheet_id": spreadsheet_id,
        "spreadsheet_url": url,
        "source_file": None,
        "sheet_name": sheet_name,
        "sheet_id": sheet_id,
        "headers": headers,
        "rows": normalized_rows,
        "total_rows": len(normalized_rows),
        "format_info": format_info
    }


def _extract_gsheets_format_info(sheet_data: dict) -> dict:
    format_info = {
        "column_widths": [],
        "frozen_rows": 0,
        "header_background_color": None,
        "header_text_bold": False,
        "header_text_color": None,
    }

    props = sheet_data.get('properties', {})
    grid_props = props.get('gridProperties', {})
    format_info["frozen_rows"] = grid_props.get('frozenRowCount', 0)

    grid_data = sheet_data.get('data', [{}])[0]
    col_metadata = grid_data.get('columnMetadata', [])
    format_info["column_widths"] = [
        col.get('pixelSize', 100) for col in col_metadata
    ]

    row_data = grid_data.get('rowData', [])
    if row_data:
        cells = row_data[0].get('values', [])
        if cells:
            fmt = cells[0].get('userEnteredFormat', {})
            bg = fmt.get('backgroundColor')
            if bg:
                format_info["header_background_color"] = bg
            text_fmt = fmt.get('textFormat', {})
            format_info["header_text_bold"] = text_fmt.get('bold', False)
            fg = text_fmt.get('foregroundColor')
            if fg:
                format_info["header_text_color"] = fg

    return format_info


# ──────────────────────────────────────────
# 공통 유틸
# ──────────────────────────────────────────

def update_run_log(phase: int, status: str, error: str | None = None):
    log_path = Path('output/run_log.json')
    from datetime import datetime

    log = {"run_id": None, "phases": [], "escalations": []}
    if log_path.exists():
        with open(log_path, encoding='utf-8') as f:
            log = json.load(f)

    if not log.get("run_id"):
        log["run_id"] = datetime.now().isoformat()

    entry = {"phase": phase, "status": status, "timestamp": datetime.now().isoformat()}
    if error:
        entry["error"] = error

    log["phases"] = [p for p in log["phases"] if p["phase"] != phase]
    log["phases"].append(entry)

    with open(log_path, 'w', encoding='utf-8') as f:
        json.dump(log, f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description='이벤트 데이터 읽기')
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--file', help='xlsx 파일 경로')
    group.add_argument('--url', help='Google Sheets URL (API 인증 필요)')
    args = parser.parse_args()

    Path('output').mkdir(exist_ok=True)

    if args.file:
        print(f"[Phase 1] xlsx 읽기 시작")
        print(f"  파일: {args.file}")
        try:
            data = read_xlsx(args.file)
        except Exception as e:
            print(f"ERROR: {e}")
            update_run_log(1, "failed", str(e))
            sys.exit(1)
    else:
        print(f"[Phase 1] Sheets 읽기 시작")
        print(f"  URL: {args.url}")
        try:
            data = read_sheet_url(args.url)
        except Exception as e:
            print(f"ERROR: {e}")
            update_run_log(1, "failed", str(e))
            sys.exit(1)

    output_path = Path('output/raw_sheet_data.json')
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    update_run_log(1, "success")
    print(f"  완료: {data['total_rows']}행 읽음")
    print(f"  헤더({len(data['headers'])}열): {data['headers']}")
    print(f"  저장: {output_path}")


if __name__ == '__main__':
    main()
